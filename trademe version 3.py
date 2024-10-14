#updates = removed redundant sheet 2 for selling info, removed redundant sold dates column from sheet 1
import os
import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from webdriver_manager.chrome import ChromeDriverManager
from openpyxl import load_workbook
import time
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from bs4 import BeautifulSoup

# Setup Chrome options for headless execution
chrome_options = Options()
chrome_options.add_argument("--headless")
chrome_options.add_argument("--no-sandbox")
chrome_options.add_argument("--disable-dev-shm-usage")

# Initialize WebDriver
driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=chrome_options)


# Function to get all product links from a store
def get_all_product_links(store_url):
    product_links = []

    # Start from page 1
    page_number = 1

    while True:
        # Construct the pagination URL
        if page_number == 1:
            current_page_url = store_url  # For the main page
        else:
            current_page_url = f"{store_url}?page={page_number}"  # For subsequent pages

        print(f"Accessing: {current_page_url}")
        driver.get(current_page_url)

        try:
            # Wait for products to load
            WebDriverWait(driver, 20).until(
                EC.presence_of_element_located((By.CSS_SELECTOR, "a.tm-marketplace-search-card__detail-section--link"))
            )
            product_elements = driver.find_elements(By.CSS_SELECTOR,
                                                    "a.tm-marketplace-search-card__detail-section--link")
            print(f"Found {len(product_elements)} products on page {page_number}.")
            for elem in product_elements:
                product_url = elem.get_attribute('href')
                product_links.append(product_url)
        except Exception as e:
            print(f"Error extracting product links: {e}")
            break

        # Handle pagination: Check if next page exists
        try:
            next_button = driver.find_element(By.CSS_SELECTOR, 'a[rel="nofollow"][aria-label*="Next page"]')
            if next_button:
                page_number += 1
            else:
                break
        except Exception as e:
            print("No more pages found")
            break

    print(f"Total product links found: {len(product_links)}")
    return product_links

# Function to get product information using Selenium
def get_product_info(url):
    driver.get(url)
    time.sleep(10)  # Ensure page is fully loaded

    # Scraping product title
    try:
        title = driver.find_element(By.CLASS_NAME, 'tm-marketplace-buyer-options__listing_title').text.strip()
    except:
        title = 'N/A'

    # Scraping listing number
    try:
        listing_number_divs = driver.find_elements(By.CLASS_NAME, 'tm-share-listing-link__info-block')
        listing_number = 'N/A'
        for div in listing_number_divs:
            text = div.text.strip()
            if 'Listing #' in text:
                listing_number = text.split('Listing #')[-1].strip()
                break
    except Exception as e:
        print(f"Error finding listing number: {e}")
        listing_number = 'N/A'

    # Scraping SKU
    try:
        script = driver.find_element(By.ID, 'frend-state').get_attribute('innerHTML')
        start_index = script.find('"sKU":"') + len('"sKU":"')
        end_index = script.find('"', start_index)
        sku = script[start_index:end_index]
    except:
        sku = 'N/A'

    # Scraping category (complete breadcrumb)
    try:
        breadcrumbs = driver.find_elements(By.CLASS_NAME, 'o-breadcrumbs__item')
        category = ' > '.join([crumb.text.strip() for crumb in breadcrumbs if crumb.text.strip()])
    except Exception as e:
        print(f"Error finding category: {e}")
        category = 'N/A'

    # Scraping product description
    try:
        description_element = driver.find_element(By.CLASS_NAME, 'tm-markdown')
        description = description_element.text.strip()
    except:
        description = 'N/A'

    # Scraping selling price
    try:
        price_element = driver.find_element(By.CSS_SELECTOR, '.tm-buy-now-box__price strong')
        price = price_element.text.strip()
    except:
        price = 'N/A'

    # Scraping quantity
    try:
        quantity_element = driver.find_element(By.CSS_SELECTOR, 'input[name="quantity"]').get_attribute("value")
        quantity = quantity_element if quantity_element else 'N/A'
    except:
        quantity = 'N/A'

    # Scraping available quantity
    try:
        available_element = driver.find_element(By.CSS_SELECTOR, 'input[name="quantityRemaining"]').get_attribute("value")
        available = available_element if available_element else 'N/A'
    except:
        available = 'N/A'

    # Scraping images using the updated approach
    images = []
    for i in range(1, 6):  # Assuming there are up to 5 images
        try:
            image_element = driver.find_element(By.CSS_SELECTOR, f'.tm-marketplace-listing-photos__thumbnail-slider-item:nth-child({i}) .o-aspect-ratio')
            image_url = image_element.get_attribute('style')
            start_index = image_url.find('url("') + len('url("')
            end_index = image_url.find('")', start_index)
            image_url = image_url[start_index:end_index]
            images.append(image_url)
        except:
            images.append('N/A')

    # Ensure that we have 5 slots for images
    images += ['N/A'] * (5 - len(images))

    # Scraping product condition
    try:
        details_element = driver.find_element(By.CSS_SELECTOR, '.o-rack-item__secondary')
        details = details_element.text.strip()
    except:
        details = 'N/A'

    # Scraping shipping options
    try:
        shipping_elements = driver.find_elements(By.CSS_SELECTOR, 'tbody tr')
        shipping_data = []
        for row in shipping_elements:
            cols = row.find_elements(By.TAG_NAME, 'td')
            if len(cols) >= 2:
                shipping_info = f"{cols[0].text.strip()}\t{cols[1].text.strip()}"
                shipping_data.append(shipping_info)
        shipping = '\n'.join(shipping_data) if shipping_data else 'N/A'
    except:
        shipping = 'N/A'

    # Scraping page views
    try:
        page_views_element = driver.find_element(By.CSS_SELECTOR, '.tm-share-listing-link__info-block b')
        page_views = page_views_element.text.strip()
    except:
        page_views = 'N/A'

    # Construct the data dictionary
    data = {
        "Link": url,
        "Listing Number": listing_number,
        "SKU": sku,
        "Category": category,
        "Product Title": title,
        "Selling Price": price,
        "Quantity": quantity,
        "Available": available,
        "Main Image": images[0],
        "Image #2": images[1],
        "Image #3": images[2],
        "Image #4": images[3],
        "Image #5": images[4],
        "Details": details,
        "Description": description,
        "Shipping": shipping,
        "Page Views": page_views,
    }

    return data

# Function to save data to an Excel file
def save_to_excel(all_data, filename="trademe_scraped_data.xlsx"):
    file_path = os.path.join("/content/", filename)
    os.makedirs(os.path.dirname(file_path), exist_ok=True)

    df = pd.DataFrame(all_data)

    if os.path.exists(file_path):
        try:
            book = load_workbook(file_path)
            writer = pd.ExcelWriter(file_path, engine='openpyxl')
            writer.book = book

            # Saving the main scraped data
            if 'Scraped Data' in writer.book.sheetnames:
                startrow = writer.book['Scraped Data'].max_row
            else:
                startrow = 0

            df.to_excel(writer, index=False, header=startrow == 0, sheet_name='Scraped Data', startrow=startrow)
            writer.close()
        except Exception as e:
            print(f"Error loading workbook: {e}")
    else:
        try:
            writer = pd.ExcelWriter(file_path, engine='openpyxl')
            df.to_excel(writer, index=False, header=True, sheet_name='Scraped Data')
            writer.close()
        except Exception as e:
            print(f"Error saving Excel file: {e}")

    print(f"File saved at: {file_path}")

# Main function
def main():
    store_url = input("Enter the URL of the TradeMe store: ").strip()

    # If the URL contains a page number, remove it to start from the main page
    if "page=" in store_url:
        store_url = store_url.split("?")[0]

    print(f"Processing store: {store_url}")
    try:
        product_links = get_all_product_links(store_url)
        print(f"Total products found: {len(product_links)}")

        all_data = []
        for idx, link in enumerate(product_links, start=1):
            try:
                product_data = get_product_info(link)
                all_data.append(product_data)
                print(f"Successfully processed data for product: {link} [{idx}/{len(product_links)}]")
            except Exception as e:
                print(f"An error occurred with product URL {link}: {e}")

        save_to_excel(all_data)
        print(f"Data saved to Excel file.")

    except Exception as e:
        print(f"An error occurred with store URL {store_url}: {e}")

# Initialize ChromeDriver and start scraping
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service, options=chrome_options)

main()
driver.quit()
